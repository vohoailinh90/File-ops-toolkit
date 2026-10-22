import io
import os
import threading
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

from openpyxl import Workbook

from file_ops_toolkit.copier import copy_without_overwrite
from file_ops_toolkit.excel import ExcelError, JobRow, read_jobs
from file_ops_toolkit.finder import FolderError, find_matches, scan_folder
from file_ops_toolkit.paths import AddressError, clean_file_name, looks_like_path, normalize_address, safe_dir_name
from file_ops_toolkit.runner import DestinationError, JobRun, Options, run_job


def touch(path: Path, content: str = "x", mtime: int | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content)
    if mtime is not None:
        os.utime(path, (mtime, mtime))
    return path


def workbook_bytes(rows, hyperlinks=None, formulas=None) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.append(["Address", "File_Name"])
    for row in rows:
        ws.append(list(row))
    for cell, target in (hyperlinks or {}).items():
        ws[cell].hyperlink = target
    for cell, formula in (formulas or {}).items():
        ws[cell] = formula
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


class AddressTests(unittest.TestCase):
    def test_unc_path_is_kept_as_is(self):
        raw = r"\\fileserver\reports\Folders_For_QMC22\19QMC_0039_230005433535"
        self.assertEqual(normalize_address(raw), raw)

    def test_file_urls_become_windows_paths(self):
        cases = {
            "file://fileserver/reports/a%20b": r"\\fileserver\reports\a b",
            "file:///C:/Reports/x": r"C:\Reports\x",
            r"file:///\\fileserver\reports\x": r"\\fileserver\reports\x",
            "file:////fileserver/reports/x": r"\\fileserver\reports\x",
            "file://fileserver/reports/x#sheet": r"\\fileserver\reports\x",
        }
        for url, expected in cases.items():
            with self.subTest(url=url):
                self.assertEqual(normalize_address(url), expected)

    def test_quotes_and_forward_slash_unc(self):
        self.assertEqual(normalize_address('  "\\\\srv\\share\\d"  '), r"\\srv\share\d")
        self.assertEqual(normalize_address("//srv/share/d"), r"\\srv\share\d")

    def test_web_links_and_blanks_are_refused_with_a_code(self):
        for raw, code in (("https://contoso.sharepoint.com/x", "unsupported_link"), ("  ", "missing_address")):
            with self.subTest(raw=raw), self.assertRaises(AddressError) as ctx:
                normalize_address(raw)
            self.assertEqual(ctx.exception.code, code)

    def test_text_that_is_not_an_absolute_path_is_refused(self):
        for raw in (".", "reports/2026", "Open folder", "..\\x"):
            with self.subTest(raw=raw), self.assertRaises(AddressError) as ctx:
                normalize_address(raw)
            self.assertEqual(ctx.exception.code, "invalid_address")

    def test_a_drive_letter_is_not_mistaken_for_a_url_scheme(self):
        self.assertEqual(normalize_address("D:\\reports"), "D:\\reports")
        self.assertTrue(looks_like_path("D:\\reports"))
        self.assertFalse(looks_like_path("Open folder"))

    def test_file_names_with_separators_or_parent_references_are_invalid(self):
        for bad in ("..\\secret.txt", "a/b.txt", "..", "a\x00b"):
            with self.subTest(bad=bad), self.assertRaises(AddressError) as ctx:
                clean_file_name(bad)
            self.assertEqual(ctx.exception.code, "invalid_name")
        self.assertEqual(clean_file_name("  report_0501.xlsx "), "report_0501.xlsx")

    def test_dir_names_are_made_safe(self):
        self.assertEqual(safe_dir_name('a:b*c?'), "a_b_c_")
        self.assertEqual(safe_dir_name("..."), "folder")


class ExcelTests(unittest.TestCase):
    def test_reads_text_cells_from_row_two_with_excel_row_numbers(self):
        rows = read_jobs(workbook_bytes([(r"\\srv\a", "r1.xlsx"), (None, None), (r"\\srv\b", "r2.xlsx")]))
        self.assertEqual([(r.row, r.address, r.file_name, r.error) for r in rows],
                         [(2, r"\\srv\a", "r1.xlsx", None), (4, r"\\srv\b", "r2.xlsx", None)])

    def test_hyperlink_target_wins_over_friendly_text(self):
        rows = read_jobs(workbook_bytes([("Open folder", "r1.xlsx")], hyperlinks={"A2": "file://srv/share/dir%201"}))
        self.assertEqual(rows[0].address, r"\\srv\share\dir 1")

    def test_hyperlink_formula_is_read(self):
        wb = workbook_bytes([("x", "r1.xlsx")], formulas={"A2": '=HYPERLINK("\\\\srv\\share\\d","open")'})
        self.assertEqual(read_jobs(wb)[0].address, r"\\srv\share\d")

    def test_invalid_rows_are_kept_with_their_reason(self):
        rows = read_jobs(workbook_bytes([("https://x.example/f", "a.xlsx"), (r"\\srv\a", None), (r"\\srv\a", "..\\x")]))
        self.assertEqual([r.error for r in rows], ["unsupported_link", "missing_name", "invalid_name"])

    def test_numeric_file_names_lose_their_decimal_point(self):
        self.assertEqual(read_jobs(workbook_bytes([(r"\\srv\a", 501.0)]))[0].file_name, "501")

    def test_a_number_formatted_with_leading_zeros_keeps_them(self):
        wb = workbook_bytes([(r"\\srv\a", 5), (r"\\srv\a", 1.5)])
        buffer = io.BytesIO()
        from openpyxl import load_workbook
        book = load_workbook(wb)
        book.active["B2"].number_format = "0000"
        book.save(buffer)
        buffer.seek(0)
        self.assertEqual([r.file_name for r in read_jobs(buffer)], ["0005", "1.5"])

    def test_not_a_workbook_and_empty_list_have_codes(self):
        with self.assertRaises(ExcelError) as ctx:
            read_jobs(io.BytesIO(b"not an xlsx"))
        self.assertEqual(ctx.exception.code, "unreadable_workbook")
        with self.assertRaises(ExcelError) as ctx:
            read_jobs(workbook_bytes([]))
        self.assertEqual(ctx.exception.code, "empty_workbook")


class FinderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_exact_name_matches_without_case_and_recurses(self):
        touch(self.root / "Report_0501.XLSX")
        touch(self.root / "sub" / "report_0501.xlsx")
        index = scan_folder(str(self.root))
        self.assertEqual(len(find_matches(index, "report_0501.xlsx")), 2)

    def test_name_without_extension_matches_any_extension_only_when_no_exact_file(self):
        touch(self.root / "Report_0501.xlsx")
        touch(self.root / "Report_0501.pdf")
        index = scan_folder(str(self.root))
        self.assertEqual(sorted(Path(p).name for p in find_matches(index, "Report_0501")),
                         ["Report_0501.pdf", "Report_0501.xlsx"])
        touch(self.root / "Report_0501")  # now an exact file exists, the stem fallback is not used
        index = scan_folder(str(self.root))
        self.assertEqual([Path(p).name for p in find_matches(index, "Report_0501")], ["Report_0501"])

    def test_dotted_name_still_falls_back_to_stem(self):
        touch(self.root / "Report_v1.2.xlsx")
        self.assertEqual(len(find_matches(scan_folder(str(self.root)), "Report_v1.2")), 1)

    def test_non_recursive_and_lock_files(self):
        touch(self.root / "a.txt")
        touch(self.root / "sub" / "a.txt")
        touch(self.root / "~$a.txt")
        self.assertEqual(len(find_matches(scan_folder(str(self.root), recursive=False), "a.txt")), 1)
        self.assertEqual(len(find_matches(scan_folder(str(self.root)), "a.txt")), 2)
        self.assertEqual(find_matches(scan_folder(str(self.root)), "~$a.txt"), [])

    def test_a_name_with_an_extension_never_matches_a_longer_name(self):
        touch(self.root / "report.pdf.bak")
        touch(self.root / "report.pdf")
        self.assertEqual([Path(p).name for p in find_matches(scan_folder(str(self.root)), "report.pdf")], ["report.pdf"])
        (self.root / "report.pdf").unlink()
        self.assertEqual(find_matches(scan_folder(str(self.root)), "report.pdf"), [])

    def test_missing_folder_and_file_in_place_of_folder(self):
        with self.assertRaises(FolderError) as ctx:
            scan_folder(str(self.root / "nope"))
        self.assertEqual(ctx.exception.code, "folder_missing")
        file = touch(self.root / "f.txt")
        with self.assertRaises(FolderError) as ctx:
            scan_folder(str(file))
        self.assertEqual(ctx.exception.code, "folder_not_a_directory")


class CopierTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_copy_keeps_content_and_leaves_the_source_alone(self):
        src = touch(self.root / "src" / "a.txt", "hello", mtime=1_700_000_000)
        out = copy_without_overwrite(str(src), str(self.root / "dest"))
        self.assertFalse(out.already_there)
        self.assertEqual(Path(out.dest).read_text(), "hello")
        self.assertTrue(src.exists())
        self.assertEqual(int(Path(out.dest).stat().st_mtime), 1_700_000_000)
        self.assertEqual(list((self.root / "dest").glob("*.part")), [])

    def test_identical_copy_from_an_earlier_run_is_reported_not_duplicated(self):
        src = touch(self.root / "src" / "a.txt", "hello", mtime=1_700_000_000)
        copy_without_overwrite(str(src), str(self.root / "dest"))
        again = copy_without_overwrite(str(src), str(self.root / "dest"))
        self.assertTrue(again.already_there)
        self.assertEqual(len(list((self.root / "dest").iterdir())), 1)

    def test_a_different_file_with_the_same_name_is_never_overwritten(self):
        touch(self.root / "dest" / "a.txt", "mine", mtime=1_600_000_000)
        src = touch(self.root / "src" / "a.txt", "theirs!", mtime=1_700_000_000)
        out = copy_without_overwrite(str(src), str(self.root / "dest"))
        self.assertEqual(Path(out.dest).name, "a (2).txt")
        self.assertEqual((self.root / "dest" / "a.txt").read_text(), "mine")

    def test_an_existing_part_file_of_the_user_is_never_deleted_or_overwritten(self):
        src = touch(self.root / "src" / "report.txt", "new")
        mine = touch(self.root / "dest" / "report.txt.part", "my unfinished data")
        out = copy_without_overwrite(str(src), str(self.root / "dest"))
        self.assertEqual(mine.read_text(), "my unfinished data")
        self.assertEqual(Path(out.dest).read_text(), "new")
        self.assertEqual(sorted(p.name for p in (self.root / "dest").iterdir()), ["report.txt", "report.txt.part"])

    def test_a_failed_copy_leaves_no_partial_file(self):
        src = touch(self.root / "src" / "a.txt", "hello")
        with mock.patch("file_ops_toolkit.copier.shutil.copyfileobj", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                copy_without_overwrite(str(src), str(self.root / "dest"))
        self.assertEqual(list((self.root / "dest").iterdir()), [])
        with self.assertRaises(OSError):
            copy_without_overwrite(str(self.root / "missing.txt"), str(self.root / "dest"))


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.dest = str(self.root / "dest")
        self.a = self.root / "srv" / "19QMC_0039"
        self.b = self.root / "srv" / "19QMC_0042"
        touch(self.a / "report_0501.xlsx", "a1")
        touch(self.a / "deep" / "report_0501.xlsx", "a2")
        touch(self.b / "report_0128.pdf", "b1")

    def jobs(self):
        return [
            JobRow(2, str(self.a), "report_0501.xlsx"),
            JobRow(3, str(self.b), "report_0128"),
            JobRow(4, str(self.b), "missing.xlsx"),
            JobRow(5, str(self.root / "gone"), "x.xlsx"),
            JobRow(6, "", "x", "unsupported_link"),
        ]

    def test_every_row_gets_a_status_and_all_matches_are_copied(self):
        results = run_job(self.jobs(), self.dest)
        self.assertEqual([(r.row, r.status) for r in results],
                         [(2, "copied"), (3, "copied"), (4, "not_found"), (5, "folder_missing"), (6, "unsupported_link")])
        self.assertEqual(len(results[0].files), 2)
        names = sorted(p.name for p in Path(self.dest).iterdir())
        self.assertEqual(names, ["report_0128.pdf", "report_0501 (2).xlsx", "report_0501.xlsx"])

    def test_second_run_over_the_same_list_changes_nothing(self):
        run_job(self.jobs(), self.dest)
        before = sorted(p.name for p in Path(self.dest).iterdir())
        results = run_job(self.jobs(), self.dest)
        self.assertEqual(results[0].status, "already_copied")
        self.assertEqual(sorted(p.name for p in Path(self.dest).iterdir()), before)

    def test_by_folder_layout_puts_each_source_in_its_own_subfolder(self):
        run_job(self.jobs()[:2], self.dest, Options(layout="by_folder"))
        self.assertTrue((Path(self.dest) / "19QMC_0039" / "report_0501.xlsx").exists())
        self.assertTrue((Path(self.dest) / "19QMC_0042" / "report_0128.pdf").exists())

    def test_a_destination_inside_or_equal_to_a_source_folder_is_refused_before_anything_is_written(self):
        for dest in (str(self.a), str(self.a / "out"), str(self.a / "deep" / "x")):
            with self.subTest(dest=dest), self.assertRaises(DestinationError) as ctx:
                run_job(self.jobs(), dest)
            self.assertEqual(ctx.exception.code, "dest_inside_source")
        self.assertFalse((self.a / "out").exists())
        run_job(self.jobs()[:1], str(self.root / "srv" / "19QMC_0039_other"))  # a sibling that shares a prefix is fine

    def test_the_by_folder_layout_cannot_aim_a_copy_at_a_source_folder(self):
        # destination = the parent of the sources: flat is fine, but by_folder would write into 19QMC_0039 itself
        parent = str(self.a.parent)
        run_job(self.jobs()[:1], parent)
        for layout in ("by_folder",):
            with self.subTest(layout=layout), self.assertRaises(DestinationError) as ctx:
                run_job(self.jobs()[:1], parent, Options(layout=layout))
            self.assertEqual(ctx.exception.code, "dest_inside_source")
        self.assertFalse((self.a / "report_0501.xlsx.part").exists())
        self.assertEqual(sorted(p.name for p in self.a.iterdir()), ["deep", "report_0501.xlsx"])

    def test_a_destination_reached_through_a_symlink_into_a_source_is_refused(self):
        alias = self.root / "alias"
        try:
            alias.symlink_to(self.a, target_is_directory=True)
        except (OSError, NotImplementedError):
            self.skipTest("this system cannot create symlinks")
        for dest in (str(alias), str(alias / "out"), str(alias / "not" / "yet" / "there")):
            with self.subTest(dest=dest), self.assertRaises(DestinationError):
                run_job(self.jobs()[:1], dest)
        self.assertEqual(sorted(p.name for p in self.a.iterdir()), ["deep", "report_0501.xlsx"])

    def test_sources_are_never_modified(self):
        before = {p: p.read_text() for p in (self.root / "srv").rglob("*") if p.is_file()}
        run_job(self.jobs(), self.dest)
        self.assertEqual({p: p.read_text() for p in (self.root / "srv").rglob("*") if p.is_file()}, before)

    def test_one_failing_copy_is_reported_and_the_rest_continue(self):
        real = copy_without_overwrite

        def flaky(src, dest_dir):
            if src.endswith("report_0128.pdf"):
                raise PermissionError(13, "Access is denied")
            return real(src, dest_dir)

        with mock.patch("file_ops_toolkit.runner.copy_without_overwrite", flaky):
            results = run_job([JobRow(2, str(self.b), "report_0128.pdf"), JobRow(3, str(self.a), "report_0501.xlsx")],
                              self.dest)
        self.assertEqual([r.status for r in results], ["error", "copied"])
        self.assertEqual(results[0].detail, "Access is denied")

    def test_cancel_stops_inside_a_folder_and_keeps_what_was_done(self):
        cancel = threading.Event()
        rows = [JobRow(2, str(self.b), "report_0128.pdf"), JobRow(3, str(self.b), "missing.xlsx"),
                JobRow(4, str(self.a), "report_0501.xlsx")]
        results = run_job(rows, self.dest, on_result=lambda r: cancel.set(), cancel=cancel)
        self.assertEqual([r.row for r in results], [2])
        self.assertEqual(sorted(p.name for p in Path(self.dest).iterdir()), ["report_0128.pdf"])

    def test_a_folder_shared_by_many_rows_is_scanned_once(self):
        rows = [JobRow(2, str(self.a), "report_0501.xlsx"), JobRow(3, str(self.b), "report_0128.pdf"),
                JobRow(4, str(self.a), "nope.xlsx"), JobRow(5, str(self.a).upper() if os.name == "nt" else str(self.a), "x")]
        with mock.patch("file_ops_toolkit.runner.scan_folder", wraps=scan_folder) as scan:
            run_job(rows, self.dest)
        self.assertEqual(scan.call_count, 2)

    def test_results_a_poller_already_holds_are_a_prefix_of_the_final_list(self):
        seen = []

        class Recording(JobRun):
            def _record(self, result):
                super()._record(result)
                seen.append([r["row"] for r in self.snapshot()["results"]])

        # rows 3 and 5 share a folder, rows 2 and 4 another: completion order differs from row order
        rows = [JobRow(3, str(self.a), "report_0501.xlsx"), JobRow(2, str(self.b), "report_0128.pdf"),
                JobRow(5, str(self.a), "nope"), JobRow(4, str(self.b), "nope")]
        run = Recording(rows, self.dest, Options())
        run.start()
        run.join(10)
        final = [r["row"] for r in run.snapshot()["results"]]
        self.assertEqual(sorted(final), [2, 3, 4, 5])
        for held in seen:
            self.assertEqual(final[:len(held)], held)
        for since in range(len(final) + 1):
            self.assertEqual([r["row"] for r in run.snapshot(since)["results"]], final[since:])

    def test_job_run_reports_progress_and_final_state(self):
        run = JobRun(self.jobs(), self.dest, Options())
        run.start()
        run.join(10)
        snap = run.snapshot()
        self.assertEqual((snap["state"], snap["done"], snap["total"]), ("done", 5, 5))
        self.assertEqual(snap["counts"]["copied"], 2)
        self.assertEqual(snap["copied_files"], 3)
        self.assertEqual(run.snapshot(since=5)["results"], [])
        with self.assertRaises(ValueError):
            run_job(self.jobs(), self.dest, Options(layout="nope"))


if __name__ == "__main__":
    unittest.main()
