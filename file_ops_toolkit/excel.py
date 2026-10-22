"""Read the job list from an .xlsx: row 1 is the header, A2 downwards is folder / file name."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from .paths import AddressError, clean_file_name, looks_like_path, normalize_address

MAX_ROWS = 50_000
_HYPERLINK_FORMULA = re.compile(r'^=\s*HYPERLINK\(\s*"((?:[^"]|"")*)"', re.IGNORECASE)


class ExcelError(Exception):
    """The workbook cannot be read. `code` is a catalog key under `errors`."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class JobRow:
    row: int  # the Excel row number, so a message can name the line the user sees
    address: str  # normalized folder path; "" when the row is invalid
    file_name: str
    error: str | None = None  # catalog key under `status` when the row cannot run


def read_jobs(source: str | Path | BinaryIO) -> list[JobRow]:
    try:
        values = load_workbook(source, data_only=True)
        if hasattr(source, "seek"):
            source.seek(0)
        formulas = load_workbook(source, data_only=False)
    except Exception as exc:  # openpyxl raises many types for a corrupt or non-xlsx file
        raise ExcelError("unreadable_workbook") from exc
    sheet, formula_sheet = values.active, formulas.active
    if sheet is None or formula_sheet is None:
        raise ExcelError("empty_workbook")
    rows: list[JobRow] = []
    for number in range(2, sheet.max_row + 1):
        if len(rows) >= MAX_ROWS:
            raise ExcelError("too_many_rows")
        cell_a, cell_b = sheet.cell(number, 1), sheet.cell(number, 2)
        raw_a = _text(cell_a.value)
        raw_b = _text(cell_b.value)
        if not raw_a and not raw_b and cell_a.hyperlink is None:
            continue  # a blank row is not a job
        rows.append(_job(number, _address_text(cell_a, formula_sheet.cell(number, 1)), raw_b))
    if not rows:
        raise ExcelError("empty_workbook")
    return rows


def _job(number: int, address_text: str, name_text: str) -> JobRow:
    try:
        address = normalize_address(address_text)
        name = clean_file_name(name_text)
    except AddressError as exc:
        return JobRow(number, "", name_text.strip(), exc.code)
    return JobRow(number, address, name)


def _address_text(cell, formula_cell) -> str:
    """The cell's visible text when it is itself a path; otherwise the hyperlink it carries."""
    text = _text(cell.value)
    if looks_like_path(text):
        return text
    target = cell.hyperlink.target if cell.hyperlink is not None and cell.hyperlink.target else ""
    if target:
        return target
    formula = formula_cell.value
    if isinstance(formula, str):
        match = _HYPERLINK_FORMULA.match(formula)
        if match:
            return match.group(1).replace('""', '"')
    return text


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
